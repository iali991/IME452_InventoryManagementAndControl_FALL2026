#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
 LOT-SIZING DECISION TOOL  (EOQ / EPQ / All-Units Quantity Discounts)
===============================================================================
 Companion industrial tool for the lecture
   "Order Quantities When Demand is Approximately Level"  -  Dr. Islam Ali

 WHAT IT DOES
   For every item (SKU) listed in an Excel input workbook it recommends
   HOW MUCH to order / produce and WHEN to reorder, using the models of the
   lecture:
     * EOQ  (Wilson formula)          - purchased items, instantaneous delivery
     * EPQ  (finite production rate)  - items produced in-house at rate m
     * All-units quantity discounts   - supplier price breaks (one or many)
   and it respects practical shop-floor constraints the textbook model ignores
   (pack/case size, minimum order quantity, maximum order quantity).

   It then compares the recommendation with the CURRENT policy, quantifies the
   savings and the Percentage Cost Penalty (PCP), computes a "flexibility band"
   (range of order sizes within a chosen cost tolerance) and writes:
     1. <name>_results.xlsx  - formatted Excel workbook with all numbers
     2. <name>_report.pdf    - management report with charts (portfolio + per SKU)

 HOW TO RUN  (see USER_GUIDE.md for the full manual)
   python lot_sizing_tool.py --template  input/my_input.xlsx     # blank template
   python lot_sizing_tool.py --input input/example_company_input.xlsx --outdir output
   python lot_sizing_tool.py --input my.xlsx --outdir out --no-pdf   # Excel only
   python lot_sizing_tool.py --input my.xlsx --sku-pages 10          # detail pages for top-10 SKUs

   or from Python / Jupyter:
     import lot_sizing_tool as lst
     res = lst.run("input/example_company_input.xlsx", "output")

 REQUIREMENTS: python >= 3.9, numpy, pandas, matplotlib, openpyxl
===============================================================================
"""
from __future__ import annotations

import argparse
import datetime as _dt
import math
import os
import sys
import textwrap
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import matplotlib

if __name__ == "__main__":
    matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.ticker import FuncFormatter

__version__ = "1.0.0"

# -----------------------------------------------------------------------------
# Visual identity (colour-blind-safe categorical order)
# -----------------------------------------------------------------------------
C_BLUE, C_ORANGE, C_AQUA, C_YELLOW = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
C_MAGENTA, C_GREEN, C_VIOLET, C_RED = "#e87ba4", "#008300", "#4a3aa7", "#e34948"
C_INK, C_INK2, C_MUTED, C_GRID = "#0b0b0b", "#52514e", "#8a8984", "#e4e3df"
C_HEADER = "#0f4c81"

plt.rcParams.update({
    "figure.facecolor": "white", "axes.facecolor": "white",
    "axes.edgecolor": C_MUTED, "axes.labelcolor": C_INK2, "axes.titleweight": "bold",
    "axes.titlesize": 11, "axes.labelsize": 9, "xtick.color": C_INK2, "ytick.color": C_INK2,
    "xtick.labelsize": 8, "ytick.labelsize": 8, "axes.grid": True, "grid.color": C_GRID,
    "grid.linewidth": 0.7, "axes.spines.top": False, "axes.spines.right": False,
    "legend.fontsize": 8, "legend.frameon": False, "font.size": 9,
})

# -----------------------------------------------------------------------------
# Input specification  (single source of truth for template, validation, docs)
# -----------------------------------------------------------------------------
SETTINGS_SPEC = [
    # key, default, description
    ("company_name", "My Company", "Name printed on the report cover"),
    ("report_title", "Lot-Sizing Recommendations", "Title printed on the report cover"),
    ("analyst", "", "Person who prepared the analysis"),
    ("currency", "$", "Currency symbol used in outputs"),
    ("days_per_year", 250, "Working days per year (used to convert lead times and cycle times)"),
    ("default_carrying_rate", 0.24, "Carrying charge r ($/$/year) used when an item has no own value"),
    ("flex_tolerance_pct", 2.0, "Cost tolerance (%) that defines the 'flexibility band' of order sizes"),
]

ITEM_COLUMNS = [
    # column, required, type, unit, description
    ("sku", True, "text", "-", "Unique item code"),
    ("description", False, "text", "-", "Item name / description"),
    ("category", False, "text", "-", "Free grouping label (e.g. Raw material, Spare part)"),
    ("supply_type", True, "text", "Purchased/Produced",
     "'Purchased' -> EOQ logic (whole lot arrives at once); 'Produced' -> EPQ logic (lot built at rate m)"),
    ("annual_demand", True, "number", "units/year", "D - expected demand per year (approximately level)"),
    ("order_cost", True, "number", "currency/order", "A - fixed cost per purchase order or per production setup"),
    ("unit_cost", True, "number", "currency/unit", "v (or v0) - base unit variable cost, before any discount"),
    ("carrying_rate", False, "number", "$/$/year", "r - leave blank to use default_carrying_rate"),
    ("production_rate", False, "number", "units/year", "m - required for 'Produced' items, must exceed annual_demand"),
    ("lead_time_days", False, "number", "working days", "Replenishment lead time L (deterministic); used for reorder point"),
    ("pack_size", False, "number", "units", "Order must be a multiple of this (case, pallet, batch); blank = 1"),
    ("min_order_qty", False, "number", "units", "Supplier / process minimum order quantity (MOQ); blank = none"),
    ("max_order_qty", False, "number", "units", "Upper limit (storage space, shelf life, cash); blank = none"),
    ("current_order_qty", False, "number", "units", "Order size used today - for the savings comparison; blank = skip"),
]

BREAK_COLUMNS = [
    ("sku", True, "text", "-", "Item code (must exist in the Items sheet)"),
    ("min_qty", True, "number", "units", "Order quantity from which this price applies (all-units)"),
    ("unit_cost", True, "number", "currency/unit", "Unit price for the WHOLE order when Q >= min_qty"),
]


# =============================================================================
# 1. CORE MODEL  (the lecture formulas)
# =============================================================================
def eoq(A: float, D: float, v: float, r: float) -> float:
    """Wilson formula  EOQ = sqrt(2AD / vr)."""
    return math.sqrt(2.0 * A * D / (v * r))


def epq(A: float, D: float, v: float, r: float, m: float) -> float:
    """EPQ = sqrt(2AD / (vr(1-D/m))) = EOQ / sqrt(1-D/m)."""
    return math.sqrt(2.0 * A * D / (v * r * (1.0 - D / m)))


def pcp(alpha: float) -> float:
    """Percentage cost penalty of ordering Q' = (1+alpha)EOQ :  50 alpha^2 / (1+alpha)."""
    return 50.0 * alpha ** 2 / (1.0 + alpha)


@dataclass
class Item:
    sku: str
    description: str
    category: str
    supply_type: str
    D: float
    A: float
    v0: float
    r: float
    m: Optional[float]
    lead_time_days: float
    pack: float
    moq: float
    qmax: float
    current_q: Optional[float]
    breaks: List[Tuple[float, float]] = field(default_factory=list)  # sorted (min_qty, price), first min_qty=0

    # ---- derived helpers --------------------------------------------------
    @property
    def k(self) -> float:
        """Carrying factor: 1 for EOQ, (1-D/m) for EPQ."""
        if self.supply_type == "Produced" and self.m:
            return 1.0 - self.D / self.m
        return 1.0

    @property
    def model(self) -> str:
        base = "EPQ" if self.supply_type == "Produced" else "EOQ"
        return base + (" + discounts" if len(self.breaks) > 1 else "")

    def price(self, Q: float) -> float:
        """All-units price: the whole order is charged at the tier Q falls in."""
        p = self.breaks[0][1]
        for q_min, price in self.breaks:
            if Q >= q_min - 1e-9:
                p = price
        return p

    def costs(self, Q: float) -> Dict[str, float]:
        """Annual cost components for order size Q (lecture TRC + purchase cost Dv)."""
        v = self.price(Q)
        ordering = self.A * self.D / Q
        holding = Q / 2.0 * self.k * v * self.r
        purchase = self.D * v
        return dict(Q=Q, unit_price=v, ordering=ordering, holding=holding,
                    relevant=ordering + holding, purchase=purchase,
                    total=ordering + holding + purchase)

    def total(self, Q: float) -> float:
        return self.costs(Q)["total"]


def _round_candidates(q: float, pack: float) -> List[float]:
    lo = math.floor(q / pack) * pack
    hi = math.ceil(q / pack) * pack
    return [x for x in {lo, hi} if x > 0]


def optimise(item: Item) -> Dict:
    """
    Generalised lecture procedure for all-units discounts (works for 1..n breaks,
    EOQ or EPQ carrying factor) plus practical constraints:

      For each price tier i with range [b_i, b_{i+1}):
        1. compute the unconstrained optimum at that tier's price
               Q_i = sqrt(2AD / (v_i r k))           (k = 1 or 1-D/m)
        2. if Q_i lies inside the tier -> candidate;
           otherwise the best point of the tier is its nearest end
           (TRC is convex inside a tier) -> candidate = boundary.
        3. clip to [MOQ, Qmax] and round to pack-size multiples (both sides).
      Evaluate TOTAL cost (incl. purchase cost Dv) of every feasible candidate
      and keep the cheapest.  With one break this reproduces the lecture's
      Case 1 / Case 2 / Case 3 logic exactly.
    """
    it = item
    lo_bound = max(it.moq, it.pack, 1e-9)
    hi_bound = it.qmax if it.qmax > 0 else math.inf
    tiers = it.breaks
    theo_candidates = []
    practical: List[float] = []
    for i, (b_i, v_i) in enumerate(tiers):
        b_next = tiers[i + 1][0] if i + 1 < len(tiers) else math.inf
        q_star = math.sqrt(2 * it.A * it.D / (v_i * it.r * it.k))
        # best continuous point inside the tier (without practical constraints)
        if q_star < b_i:
            q_tier = b_i
        elif q_star >= b_next:
            q_tier = b_next * (1 - 1e-9)  # just below next break (dominated anyway)
        else:
            q_tier = q_star
        if q_tier > 0:
            theo_candidates.append(q_tier)
        # practical candidates
        for q in (q_tier, b_i if b_i > 0 else None):
            if q is None:
                continue
            qc = min(max(q, lo_bound), hi_bound)
            practical.extend(_round_candidates(qc, it.pack))
    # the constraint boundaries themselves are candidates too
    practical.extend(_round_candidates(lo_bound, it.pack))
    if math.isfinite(hi_bound):
        practical.append(math.floor(hi_bound / it.pack) * it.pack)

    def feasible(q):
        return (q >= it.moq - 1e-9 and q <= hi_bound + 1e-9 and q > 0
                and abs(q / it.pack - round(q / it.pack)) < 1e-6)

    practical = sorted({round(q, 6) for q in practical if feasible(q)})
    if not practical:  # e.g. qmax < moq  -> fall back to MOQ
        practical = [math.ceil(max(it.moq, it.pack) / it.pack) * it.pack]
    best_q = min(practical, key=it.total)

    theo_q = min(theo_candidates, key=it.total)
    # basic (no discount, no constraint) textbook value at base price
    base_q = math.sqrt(2 * it.A * it.D / (tiers[0][1] * it.r * it.k))

    # --- lecture "case" classification (meaningful for a single break) -----
    case = "-"
    if len(tiers) == 2:
        Q1, vd = tiers[1]
        q_d = math.sqrt(2 * it.A * it.D / (vd * it.r * it.k))
        if q_d >= Q1:
            case = "Case 3 (EOQ at discount price)"
        elif it.total(Q1) < it.total(min(base_q, Q1 * (1 - 1e-9))):
            case = "Case 1 (order break qty Q1)"
        else:
            case = "Case 2 (ignore discount)"
    elif len(tiers) > 2:
        tier_idx = sum(1 for b, _ in tiers if theo_q >= b - 1e-9)
        on_break = any(abs(theo_q - b) < 1e-6 for b, _ in tiers[1:])
        case = f"Tier {tier_idx} of {len(tiers)}" + (" (at break qty)" if on_break else " (interior optimum)")

    binding = []
    if best_q <= it.moq + 1e-9 and it.moq > 0 and theo_q < it.moq:
        binding.append("MOQ")
    if math.isfinite(hi_bound) and best_q >= math.floor(hi_bound / it.pack) * it.pack - 1e-9 and theo_q > hi_bound:
        binding.append("Max qty")
    if it.pack > 1:
        binding.append(f"Pack {it.pack:g}")
    return dict(best_q=best_q, theo_q=theo_q, base_q=base_q, case=case, binding=", ".join(binding) or "none")


def flexibility_band(item: Item, q_opt: float, tol_pct: float) -> Tuple[float, float]:
    """Contiguous range of Q around q_opt whose total cost is within tol% of the optimum
    (relevant cost basis).  Uses the lecture PCP idea numerically so it also works
    with discounts, EPQ and constraints."""
    c_opt = item.costs(q_opt)["relevant"]
    target = c_opt * (1 + tol_pct / 100.0)
    tot_opt = item.total(q_opt)
    grid_lo = np.linspace(max(q_opt * 0.2, 1e-6), q_opt, 800)[::-1]
    grid_hi = np.linspace(q_opt, q_opt * 5, 1600)
    lo = q_opt
    for q in grid_lo:
        if item.total(q) - tot_opt > target - c_opt:
            break
        lo = q
    hi = q_opt
    for q in grid_hi:
        if item.total(q) - tot_opt > target - c_opt:
            break
        hi = q
    return lo, hi


# =============================================================================
# 2. INPUT
# =============================================================================
def _num(x, default=None):
    try:
        if x is None or (isinstance(x, str) and x.strip() == ""):
            return default
        f = float(x)
        return default if math.isnan(f) else f
    except (TypeError, ValueError):
        return default


def load_input(path: str):
    """Read Settings, Items and PriceBreaks sheets."""
    xl = pd.ExcelFile(path)
    settings = {k: d for k, d, _ in SETTINGS_SPEC}
    if "Settings" in xl.sheet_names:
        s = pd.read_excel(xl, sheet_name="Settings")
        s.columns = [str(c).strip().lower() for c in s.columns]
        for _, row in s.iterrows():
            key = str(row.get("setting", "")).strip()
            if key in settings and not pd.isna(row.get("value")):
                settings[key] = row["value"]
    items = pd.read_excel(xl, sheet_name="Items")
    items.columns = [str(c).strip().lower() for c in items.columns]
    items = items.dropna(how="all")
    breaks = pd.DataFrame(columns=["sku", "min_qty", "unit_cost"])
    if "PriceBreaks" in xl.sheet_names:
        breaks = pd.read_excel(xl, sheet_name="PriceBreaks")
        breaks.columns = [str(c).strip().lower() for c in breaks.columns]
        breaks = breaks.dropna(how="all")
    for key in ("days_per_year", "default_carrying_rate", "flex_tolerance_pct"):
        settings[key] = float(settings[key])
    return settings, items, breaks


def build_items(settings, items_df, breaks_df) -> Tuple[List[Item], pd.DataFrame]:
    """Validate rows and build Item objects.  Returns (items, validation_log)."""
    log = []

    def note(sku, level, msg):
        log.append(dict(sku=sku, level=level, message=msg))

    missing_cols = [c for c, req, *_ in ITEM_COLUMNS if req and c not in items_df.columns]
    if missing_cols:
        raise ValueError(f"Items sheet is missing required column(s): {missing_cols}")

    out: List[Item] = []
    seen = set()
    for idx, row in items_df.iterrows():
        sku = str(row.get("sku", "")).strip()
        if not sku or sku.lower() == "nan":
            note(f"row {idx + 2}", "ERROR", "Empty SKU - row skipped")
            continue
        if sku in seen:
            note(sku, "ERROR", "Duplicate SKU - second occurrence skipped")
            continue
        seen.add(sku)
        st = str(row.get("supply_type", "Purchased")).strip().capitalize()
        if st not in ("Purchased", "Produced"):
            note(sku, "WARNING", f"supply_type '{st}' not recognised - treated as Purchased")
            st = "Purchased"
        D, A, v0 = _num(row.get("annual_demand")), _num(row.get("order_cost")), _num(row.get("unit_cost"))
        bad = [n for n, x in (("annual_demand", D), ("order_cost", A), ("unit_cost", v0)) if x is None or x <= 0]
        if bad:
            note(sku, "ERROR", f"Missing or non-positive {bad} - item skipped")
            continue
        r = _num(row.get("carrying_rate"), settings["default_carrying_rate"])
        if r <= 0:
            note(sku, "ERROR", "carrying_rate must be > 0 - item skipped")
            continue
        if r > 1:
            note(sku, "WARNING", f"carrying_rate {r} > 1 - did you enter a percentage? Interpreted as {r / 100:.2f}")
            r = r / 100
        m = _num(row.get("production_rate"))
        if st == "Produced":
            if m is None:
                note(sku, "WARNING", "Produced item without production_rate - EOQ (instantaneous) used")
            elif m <= D:
                note(sku, "ERROR", f"production_rate ({m:g}) must exceed annual_demand ({D:g}) - item skipped")
                continue
        L = _num(row.get("lead_time_days"), 0.0)
        pack = _num(row.get("pack_size"), 1.0) or 1.0
        moq = _num(row.get("min_order_qty"), 0.0)
        qmax = _num(row.get("max_order_qty"), 0.0)
        if qmax and qmax < max(moq, pack):
            note(sku, "WARNING", "max_order_qty is below MOQ / pack size - constraint ignored")
            qmax = 0.0
        cur = _num(row.get("current_order_qty"))
        # price schedule (all-units)
        sched = [(0.0, v0)]
        if len(breaks_df):
            b = breaks_df[breaks_df["sku"].astype(str).str.strip() == sku]
            for _, br in b.iterrows():
                q, p = _num(br.get("min_qty")), _num(br.get("unit_cost"))
                if q is None or p is None or q <= 0 or p <= 0:
                    note(sku, "WARNING", f"Invalid price break ({br.get('min_qty')}, {br.get('unit_cost')}) ignored")
                    continue
                sched.append((q, p))
        sched = sorted(sched)
        clean = [sched[0]]
        for q, p in sched[1:]:
            if p >= clean[-1][1]:
                note(sku, "WARNING", f"Break at {q:g} has price {p} not lower than previous tier - ignored")
                continue
            clean.append((q, p))
        if D / 12 < 1:
            note(sku, "INFO", "Very low demand (< 1 unit/month) - EOQ assumptions of level demand may not hold")
        out.append(Item(sku, str(row.get("description", "") or ""), str(row.get("category", "") or ""),
                        st, D, A, v0, r, m if st == "Produced" else None, L, pack, moq, qmax, cur, clean))
    return out, pd.DataFrame(log, columns=["sku", "level", "message"])


# =============================================================================
# 3. ANALYSIS
# =============================================================================
def analyse(settings, items: List[Item]) -> pd.DataFrame:
    rows = []
    dpy = settings["days_per_year"]
    tol = settings["flex_tolerance_pct"]
    for it in items:
        opt = optimise(it)
        q = opt["best_q"]
        c = it.costs(q)
        lo, hi = flexibility_band(it, q, tol)
        cycle_years = q / it.D
        max_inv = q * it.k
        row = dict(
            sku=it.sku, description=it.description, category=it.category, supply_type=it.supply_type,
            model=it.model, annual_demand=it.D, order_cost=it.A, base_unit_cost=it.v0, carrying_rate=it.r,
            production_rate=it.m, n_price_tiers=len(it.breaks),
            textbook_q=opt["base_q"], theoretical_opt_q=opt["theo_q"], recommended_q=q,
            discount_case=opt["case"], binding_constraints=opt["binding"],
            unit_price_paid=c["unit_price"], orders_per_year=it.D / q,
            cycle_time_days=cycle_years * dpy, cycle_time_weeks=cycle_years * 52,
            months_of_supply=cycle_years * 12,
            production_run_days=(q / it.m * dpy) if it.m else np.nan,
            reorder_point=it.D * it.lead_time_days / dpy, lead_time_days=it.lead_time_days,
            max_inventory=max_inv, avg_inventory=max_inv / 2, avg_inventory_value=max_inv / 2 * c["unit_price"],
            turnover_ratio=it.D / (max_inv / 2),
            annual_ordering_cost=c["ordering"], annual_holding_cost=c["holding"],
            annual_relevant_cost=c["relevant"], annual_purchase_cost=c["purchase"], annual_total_cost=c["total"],
            flex_q_low=lo, flex_q_high=hi,
        )
        if it.current_q and it.current_q > 0:
            cc = it.costs(it.current_q)
            row.update(
                current_q=it.current_q, current_orders_per_year=it.D / it.current_q,
                current_avg_inventory_value=it.current_q * it.k / 2 * cc["unit_price"],
                current_total_cost=cc["total"], annual_savings=cc["total"] - c["total"],
                current_pcp_pct=(cc["total"] - c["total"]) / c["relevant"] * 100,
            )
        else:
            row.update(current_q=np.nan, current_orders_per_year=np.nan, current_avg_inventory_value=np.nan,
                       current_total_cost=np.nan, annual_savings=np.nan, current_pcp_pct=np.nan)
        rows.append(row)
    return pd.DataFrame(rows)


def portfolio_kpis(res: pd.DataFrame) -> Dict[str, float]:
    has_cur = res["current_q"].notna()
    return dict(
        n_items=len(res),
        n_with_current=int(has_cur.sum()),
        total_cost_recommended=res["annual_total_cost"].sum(),
        relevant_cost_recommended=res["annual_relevant_cost"].sum(),
        total_cost_current=res.loc[has_cur, "current_total_cost"].sum(),
        total_savings=res.loc[has_cur, "annual_savings"].sum(),
        inv_value_recommended=res["avg_inventory_value"].sum(),
        inv_value_current=res.loc[has_cur, "current_avg_inventory_value"].sum(),
        orders_recommended=res["orders_per_year"].sum(),
        orders_current=res.loc[has_cur, "current_orders_per_year"].sum(),
        n_discount_taken=int((res["unit_price_paid"] < res["base_unit_cost"] - 1e-9).sum()),
    )


# =============================================================================
# 4. EXCEL OUTPUT
# =============================================================================
RESULT_COLUMNS = [
    # column, header, number format, explanation
    ("sku", "SKU", None, "Item code"),
    ("description", "Description", None, ""),
    ("model", "Model", None, "EOQ (purchased) or EPQ (produced); '+ discounts' when price breaks exist"),
    ("recommended_q", "Recommended order qty", "#,##0", "Q to order/produce each time (after pack/MOQ/max rules)"),
    ("theoretical_opt_q", "Theoretical optimum Q", "#,##0.0", "Continuous optimum before practical rounding / limits"),
    ("textbook_q", "Textbook EOQ/EPQ (base price)", "#,##0.0", "Wilson/EPQ formula at the base price - lecture value"),
    ("discount_case", "Discount case", None, "Lecture case 1/2/3 (single break) or tier chosen (multi-break)"),
    ("binding_constraints", "Practical rules applied", None, "Constraints that shaped the recommendation"),
    ("unit_price_paid", "Unit price paid", "#,##0.000", "All-units price at the recommended Q"),
    ("orders_per_year", "Orders (runs) per year", "#,##0.0", "D/Q"),
    ("cycle_time_days", "Cycle time (working days)", "#,##0.0", "Q/D x days per year"),
    ("months_of_supply", "Months of supply (T_EOQ)", "#,##0.00", "12 Q/D"),
    ("production_run_days", "Production run length (days)", "#,##0.0", "Q/m x days per year (EPQ items)"),
    ("reorder_point", "Reorder point (units)", "#,##0", "Reorder when inventory position = D x L"),
    ("max_inventory", "Max inventory (units)", "#,##0", "Q (EOQ) or Q(1-D/m) (EPQ)"),
    ("avg_inventory", "Average inventory (units)", "#,##0", "Max inventory / 2"),
    ("avg_inventory_value", "Average inventory value", "#,##0", "Average inventory x unit price"),
    ("turnover_ratio", "Turnover ratio", "#,##0.0", "D / average inventory"),
    ("annual_ordering_cost", "Annual ordering/setup cost", "#,##0.00", "AD/Q"),
    ("annual_holding_cost", "Annual carrying cost", "#,##0.00", "Q/2 x k x v x r"),
    ("annual_relevant_cost", "TRC (ordering + carrying)", "#,##0.00", "Lecture TRC(Q)"),
    ("annual_purchase_cost", "Annual purchase cost", "#,##0.00", "D x v"),
    ("annual_total_cost", "Total annual cost", "#,##0.00", "TRC + purchase cost"),
    ("flex_q_low", "Flex band: low Q", "#,##0", "Smallest Q within the cost tolerance"),
    ("flex_q_high", "Flex band: high Q", "#,##0", "Largest Q within the cost tolerance"),
    ("current_q", "Current order qty", "#,##0", "From input"),
    ("current_total_cost", "Current total annual cost", "#,##0.00", "Cost of today's policy"),
    ("annual_savings", "Annual savings", "#,##0.00", "Current total cost - recommended total cost"),
    ("current_pcp_pct", "Current policy PCP (%)", "0.0", "Extra cost of the current policy, % of optimal TRC"),
]


def write_excel(path, settings, res, log, items_df, breaks_df):
    kp = portfolio_kpis(res)
    cur = settings["currency"]
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        # README
        readme = [
            ["LOT-SIZING RESULTS", ""],
            ["Company", settings["company_name"]],
            ["Generated", _dt.datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["Tool version", __version__],
            ["", ""],
            ["Sheet", "Content"],
            ["Summary", "Portfolio KPIs: cost, savings, inventory investment, number of orders"],
            ["Recommendations", "One row per SKU: what to order, when to reorder, costs, savings"],
            ["Column_Guide", "Meaning and formula of every column in Recommendations"],
            ["Validation_Log", "Errors/warnings found in the input (items with ERROR were skipped)"],
            ["Input_Items / Input_PriceBreaks", "Copy of the data that was analysed (audit trail)"],
            ["", ""],
            ["How to act on it", "1) Set the order quantity in your ERP to 'Recommended order qty'."],
            ["", "2) Set the reorder point to 'Reorder point (units)' (inventory position incl. open orders)."],
            ["", "3) Any Q inside the flexibility band costs at most tol% more - use it to round to pallets/trucks."],
            ["", "4) Re-run when demand, prices or costs change by more than ~20% (EOQ is robust to smaller errors)."],
        ]
        pd.DataFrame(readme).to_excel(xw, sheet_name="README", index=False, header=False)
        summ = pd.DataFrame([
            ["Items analysed", kp["n_items"], ""],
            ["Items with a current policy", kp["n_with_current"], ""],
            ["Items where a quantity discount is taken", kp["n_discount_taken"], ""],
            [f"Total annual cost - recommended ({cur})", kp["total_cost_recommended"], "all items"],
            [f"Ordering + carrying cost (TRC) - recommended ({cur})", kp["relevant_cost_recommended"], "all items"],
            [f"Total annual cost - current ({cur})", kp["total_cost_current"], "items with current policy"],
            [f"Annual savings ({cur})", kp["total_savings"], "items with current policy"],
            [f"Average inventory value - recommended ({cur})", kp["inv_value_recommended"], "all items"],
            [f"Average inventory value - current ({cur})", kp["inv_value_current"], "items with current policy"],
            ["Orders per year - recommended", kp["orders_recommended"], "all items"],
            ["Orders per year - current", kp["orders_current"], "items with current policy"],
            ["Flexibility tolerance (%)", settings["flex_tolerance_pct"], "setting"],
            ["Working days per year", settings["days_per_year"], "setting"],
        ], columns=["KPI", "Value", "Scope"])
        summ.to_excel(xw, sheet_name="Summary", index=False)
        cols = [c for c, *_ in RESULT_COLUMNS]
        out = res[cols].copy()
        out.columns = [h for _, h, *_ in RESULT_COLUMNS]
        out.to_excel(xw, sheet_name="Recommendations", index=False)
        pd.DataFrame([(h, e) for _, h, _, e in RESULT_COLUMNS], columns=["Column", "Meaning / formula"]) \
            .to_excel(xw, sheet_name="Column_Guide", index=False)
        (log if len(log) else pd.DataFrame([["-", "OK", "No issues found"]], columns=log.columns)) \
            .to_excel(xw, sheet_name="Validation_Log", index=False)
        items_df.to_excel(xw, sheet_name="Input_Items", index=False)
        breaks_df.to_excel(xw, sheet_name="Input_PriceBreaks", index=False)
    _format_workbook(path)


def _format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    wb = load_workbook(path)
    hdr_fill = PatternFill("solid", fgColor=C_HEADER.strip("#"))
    hdr_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9D9D9")
    fmts = {h: f for _, h, f, _ in RESULT_COLUMNS}
    for ws in wb.worksheets:
        if ws.title == "README":
            ws["A1"].font = Font(bold=True, size=14, color=C_HEADER.strip("#"))
            ws["A6"].font = ws["B6"].font = Font(bold=True)
            ws.column_dimensions["A"].width = 34
            ws.column_dimensions["B"].width = 100
            continue
        for cell in ws[1]:
            cell.fill, cell.font = hdr_fill, hdr_font
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[1].height = 42
        ws.freeze_panes = "B2" if ws.title == "Recommendations" else "A2"
        for col in ws.columns:
            head = col[0].value
            width = max(10, min(46, max(len(str(c.value)) if c.value is not None else 0 for c in col[1:60]) + 2
                        if len(col) > 1 else 12))
            ws.column_dimensions[col[0].column_letter].width = max(width, 14)
            for c in col[1:]:
                c.border = Border(bottom=thin)
                if ws.title == "Recommendations" and fmts.get(head):
                    c.number_format = fmts[head]
                if ws.title == "Summary" and head == "Value":
                    c.number_format = "#,##0.00"
        if ws.title == "Recommendations":
            ws.auto_filter.ref = ws.dimensions
            # highlight savings
            green = PatternFill("solid", fgColor="E3F4EC")
            for c in ws[1]:
                if c.value in ("Recommended order qty", "Reorder point (units)", "Annual savings"):
                    for cell in ws[c.column_letter][1:]:
                        cell.fill = green
                        cell.font = Font(bold=True)
        if ws.title == "Validation_Log":
            red, amber = PatternFill("solid", fgColor="FBE3E3"), PatternFill("solid", fgColor="FDF1D6")
            for row in ws.iter_rows(min_row=2):
                if row[1].value == "ERROR":
                    for c in row:
                        c.fill = red
                elif row[1].value == "WARNING":
                    for c in row:
                        c.fill = amber
    wb.save(path)


# =============================================================================
# 5. PDF REPORT
# =============================================================================
def _money(cur):
    return FuncFormatter(lambda x, _: f"{cur}{x:,.0f}")


def _text_page(pdf, title, blocks, subtitle=None):
    fig = plt.figure(figsize=(8.27, 11.69))
    fig.text(0.07, 0.95, title, fontsize=17, weight="bold", color=C_HEADER)
    if subtitle:
        fig.text(0.07, 0.925, subtitle, fontsize=9.5, color=C_INK2)
    y = 0.89
    for head, body in blocks:
        if head:
            fig.text(0.07, y, head, fontsize=11.5, weight="bold", color=C_INK)
            y -= 0.022
        for para in body.split("\n"):
            for line in textwrap.wrap(para, 105) or [""]:
                fig.text(0.07, y, line, fontsize=9, color=C_INK2, family="DejaVu Sans")
                y -= 0.0165
        y -= 0.012
    pdf.savefig(fig)
    plt.close(fig)


def _table_pages(pdf, df, title, col_widths=None, rows_per_page=32):
    for start in range(0, len(df), rows_per_page):
        chunk = df.iloc[start:start + rows_per_page]
        fig, ax = plt.subplots(figsize=(11.69, 8.27))
        ax.axis("off")
        ax.set_title(title + (f"  (cont.)" if start else ""), loc="left", fontsize=14, color=C_HEADER, pad=12)
        tb = ax.table(cellText=chunk.values, colLabels=chunk.columns, loc="upper center",
                      cellLoc="center", colWidths=col_widths)
        tb.auto_set_font_size(False)
        tb.set_fontsize(7.2)
        tb.scale(1, 1.32)
        for (r, c), cell in tb.get_celld().items():
            cell.set_edgecolor(C_GRID)
            if r == 0:
                cell.set_facecolor(C_HEADER)
                cell.set_text_props(color="white", weight="bold")
                cell.set_height(cell.get_height() * 1.8)
            elif r % 2 == 0:
                cell.set_facecolor("#f5f7fa")
        pdf.savefig(fig)
        plt.close(fig)


def _cover(pdf, settings, kp):
    cur = settings["currency"]
    fig = plt.figure(figsize=(8.27, 11.69))
    fig.patches.append(plt.Rectangle((0, 0.72), 1, 0.28, transform=fig.transFigure, color=C_HEADER))
    fig.text(0.07, 0.90, settings["report_title"], fontsize=24, color="white", weight="bold")
    fig.text(0.07, 0.86, settings["company_name"], fontsize=15, color="white")
    fig.text(0.07, 0.76, f"Generated {_dt.date.today():%d %B %Y}"
             + (f"   |   Prepared by {settings['analyst']}" if settings.get("analyst") else "")
             + f"   |   Lot-Sizing Tool v{__version__}", fontsize=9, color="#dce8f5")
    tiles = [
        ("Items analysed", f"{kp['n_items']}"),
        ("Annual savings vs current", f"{cur}{kp['total_savings']:,.0f}"),
        ("Avg inventory value (recommended)", f"{cur}{kp['inv_value_recommended']:,.0f}"),
        ("Avg inventory value (current)", f"{cur}{kp['inv_value_current']:,.0f}"),
        ("Orders per year (recommended)", f"{kp['orders_recommended']:,.0f}"),
        ("Orders per year (current)", f"{kp['orders_current']:,.0f}"),
    ]
    for i, (lab, val) in enumerate(tiles):
        x = 0.07 + (i % 2) * 0.45
        y = 0.60 - (i // 2) * 0.14
        fig.patches.append(plt.Rectangle((x, y - 0.02), 0.41, 0.115, transform=fig.transFigure,
                                         facecolor="#f3f6fa", edgecolor=C_GRID))
        fig.text(x + 0.02, y + 0.06, lab, fontsize=9.5, color=C_INK2)
        fig.text(x + 0.02, y + 0.005, val, fontsize=20, color=C_INK, weight="bold")
    fig.text(0.07, 0.14, "Savings, current cost and current inventory refer only to items for which a current "
             "order quantity was supplied.", fontsize=8, color=C_MUTED)
    fig.text(0.07, 0.12, "Models: EOQ (Wilson), EPQ (finite production rate), all-units quantity discounts. "
             "Deterministic, level demand.", fontsize=8, color=C_MUTED)
    pdf.savefig(fig)
    plt.close(fig)


def _portfolio_charts(pdf, settings, res):
    cur = settings["currency"]
    r = res.copy()
    has = r["current_q"].notna()
    # --- page 1: savings + PCP --------------------------------------------
    fig, axes = plt.subplots(2, 1, figsize=(8.27, 11.69), gridspec_kw=dict(hspace=0.35))
    s = r[has].sort_values("annual_savings", ascending=True)
    ax = axes[0]
    ax.barh(s["sku"], s["annual_savings"], color=C_BLUE, height=0.6)
    share = s["annual_savings"].sum()
    ax.set_title(f"Annual savings by SKU (current policy -> recommended), total {cur}{share:,.0f}", loc="left")
    ax.xaxis.set_major_formatter(_money(cur))
    ax.grid(axis="y", visible=False)
    for y_, v in zip(range(len(s)), s["annual_savings"]):
        ax.text(v, y_, f" {cur}{v:,.0f}", va="center", fontsize=7, color=C_INK2)
    ax.tick_params(axis="y", labelsize=7)
    # PCP of the current policy (same SKU order as savings, largest first)
    ax = axes[1]
    p = r[has].sort_values("annual_savings", ascending=False)
    ax.bar(p["sku"], p["current_pcp_pct"], color=C_ORANGE, width=0.6, label="Current policy PCP (%)")
    ax.axhline(settings["flex_tolerance_pct"], color=C_INK2, ls="--", lw=1,
               label=f"Tolerance {settings['flex_tolerance_pct']:g}% (inside = acceptable)")
    ax.set_ylabel("Percentage cost penalty (%)")
    ax.set_title("How far is each current policy from optimal? (PCP, SKUs ordered by savings)", loc="left")
    ax.tick_params(axis="x", rotation=90, labelsize=7)
    ax.grid(axis="x", visible=False)
    ax.legend(loc="upper right")
    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)

    # --- page 2: inventory vs orders & cycle time -----------------------------
    fig, axes = plt.subplots(2, 2, figsize=(11.69, 8.27), gridspec_kw=dict(hspace=0.4, wspace=0.28))
    ax = axes[0, 0]
    ax.scatter(r.loc[has, "current_orders_per_year"], r.loc[has, "current_avg_inventory_value"], s=30,
               color=C_ORANGE, label="Current", zorder=3, edgecolor="white", lw=0.8)
    ax.scatter(r["orders_per_year"], r["avg_inventory_value"], s=30, color=C_BLUE, label="Recommended",
               zorder=4, edgecolor="white", lw=0.8)
    for _, row in r[has].iterrows():
        ax.annotate("", xy=(row["orders_per_year"], row["avg_inventory_value"]),
                    xytext=(row["current_orders_per_year"], row["current_avg_inventory_value"]),
                    arrowprops=dict(arrowstyle="->", color=C_MUTED, lw=0.7))
    ax.set_xscale("log")
    ax.set_yscale("log")
    ax.set_xlabel("Orders per year (log)")
    ax.set_ylabel(f"Average inventory value ({cur}, log)")
    ax.set_title("Trade-off: order frequency vs inventory", loc="left")
    ax.legend()
    ax = axes[0, 1]
    tot = [r.loc[has, "current_avg_inventory_value"].sum(), r.loc[has, "avg_inventory_value"].sum()]
    ords = [r.loc[has, "current_orders_per_year"].sum(), r.loc[has, "orders_per_year"].sum()]
    bars = ax.bar(["Current", "Recommended"], tot, color=[C_ORANGE, C_BLUE], width=0.5)
    for b, v, o in zip(bars, tot, ords):
        ax.text(b.get_x() + b.get_width() / 2, v, f"{cur}{v:,.0f}\n{o:,.0f} orders/yr", ha="center",
                va="bottom", fontsize=8)
    ax.set_title("Average inventory investment (items with current policy)", loc="left")
    ax.yaxis.set_major_formatter(_money(cur))
    ax.set_ylim(0, max(tot) * 1.25)
    ax.grid(axis="x", visible=False)
    ax = axes[1, 0]
    ax.hist(r["months_of_supply"], bins=min(15, max(5, len(r) // 2)), color=C_BLUE, edgecolor="white")
    ax.set_xlabel("Months of supply per order (T = 12Q/D)")
    ax.set_ylabel("Number of SKUs")
    ax.set_title("Distribution of recommended cycle times", loc="left")
    ax.grid(axis="x", visible=False)
    ax = axes[1, 1]
    comp = r.groupby("model")[["annual_ordering_cost", "annual_holding_cost"]].sum()
    x = np.arange(len(comp))
    ax.bar(x - 0.18, comp["annual_ordering_cost"], 0.34, color=C_BLUE, label="Ordering / setup")
    ax.bar(x + 0.18, comp["annual_holding_cost"], 0.34, color=C_ORANGE, label="Carrying")
    ax.set_xticks(x, comp.index, fontsize=8)
    ax.yaxis.set_major_formatter(_money(cur))
    ax.set_title("Recommended TRC components by model", loc="left")
    ax.legend()
    ax.grid(axis="x", visible=False)
    pdf.savefig(fig, bbox_inches="tight")
    plt.close(fig)


def sku_figure(it: Item, row: pd.Series, settings) -> plt.Figure:
    """One-page visual for a single SKU (also used by the lecture notebook)."""
    cur = settings["currency"]
    dpy = settings["days_per_year"]
    q = row["recommended_q"]
    fig = plt.figure(figsize=(11.69, 8.27))
    gs = fig.add_gridspec(2, 3, hspace=0.42, wspace=0.34, width_ratios=[1.3, 1.0, 0.9])
    fig.suptitle(f"{it.sku} - {it.description}   [{it.model}]", x=0.06, ha="left", fontsize=14,
                 weight="bold", color=C_HEADER)

    # (a) cost curve with tiers
    ax = fig.add_subplot(gs[0, :2])
    qmax_plot = max(q, row["theoretical_opt_q"], row["current_q"] if not pd.isna(row["current_q"]) else 0,
                    *(b for b, _ in it.breaks)) * 1.8
    grid = np.linspace(max(qmax_plot / 400, 1e-3), qmax_plot, 1200)
    for i, (b, v) in enumerate(it.breaks):
        b_next = it.breaks[i + 1][0] if i + 1 < len(it.breaks) else np.inf
        full = it.A * it.D / grid + grid / 2 * it.k * v * it.r + it.D * v
        valid = (grid >= b) & (grid < b_next)
        col = [C_BLUE, C_AQUA, C_VIOLET, C_MAGENTA, C_GREEN][i % 5]
        lab = f"price {cur}{v:g}" + (f" (Q>={b:g})" if b > 0 else "")
        ax.plot(grid, np.where(valid, np.nan, full), color=col, lw=1, ls=":", alpha=0.8)
        ax.plot(grid, np.where(valid, full, np.nan), color=col, lw=2.2, label=lab)
        if b > 0:
            ax.axvline(b, color=C_MUTED, lw=0.8, ls="--")
    ax.scatter([q], [it.total(q)], s=70, color=C_GREEN, zorder=5, label=f"Recommended Q = {q:,.0f}")
    if not pd.isna(row["current_q"]):
        ax.scatter([row["current_q"]], [it.total(row["current_q"])], s=70, marker="X", color=C_RED, zorder=5,
                   label=f"Current Q = {row['current_q']:,.0f}")
    ax.axvspan(row["flex_q_low"], row["flex_q_high"], color=C_GREEN, alpha=0.08,
               label=f"Flex band (+{settings['flex_tolerance_pct']:g}%)")
    y_ref, rel = it.total(q), it.costs(q)["relevant"]
    y_top = y_ref + 1.5 * rel
    if not pd.isna(row["current_q"]):
        y_top = max(y_top, it.total(row["current_q"]) + 0.3 * rel)
    ax.set_ylim(y_ref - 0.12 * (y_top - y_ref), y_top)
    ax.set_xlim(0, qmax_plot)
    ax.set_xlabel("Order quantity Q (units)")
    ax.set_ylabel(f"Total annual cost ({cur}/yr)")
    ax.set_title("Total annual cost curve (solid = applicable price, dotted = extension)", loc="left")
    ax.yaxis.set_major_formatter(_money(cur))
    ax.legend(fontsize=7, loc="upper right")

    # (b) key figures table
    axt = fig.add_subplot(gs[:, 2])
    axt.axis("off")
    kv = [
        ("Annual demand D", f"{it.D:,.0f}"), ("Order/setup cost A", f"{cur}{it.A:,.2f}"),
        ("Base unit cost v0", f"{cur}{it.v0:,.3f}"), ("Carrying rate r", f"{it.r:.2f}"),
        ("Production rate m", f"{it.m:,.0f}" if it.m else "-"),
        ("Textbook EOQ/EPQ", f"{row['textbook_q']:,.1f}"),
        ("Recommended Q", f"{q:,.0f}"), ("Price paid", f"{cur}{row['unit_price_paid']:,.3f}"),
        ("Orders per year", f"{row['orders_per_year']:,.1f}"),
        ("Cycle (working days)", f"{row['cycle_time_days']:,.1f}"),
        ("Reorder point", f"{row['reorder_point']:,.0f}"),
        ("TRC (order+carry)", f"{cur}{row['annual_relevant_cost']:,.0f}"),
        ("Total annual cost", f"{cur}{row['annual_total_cost']:,.0f}"),
        ("Case / tier", str(row["discount_case"]).split(" (")[0]),
        ("Rules applied", str(row["binding_constraints"])[:22]),
    ]
    if not pd.isna(row["current_q"]):
        kv += [("Current Q", f"{row['current_q']:,.0f}"), ("Annual savings", f"{cur}{row['annual_savings']:,.0f}"),
               ("Current PCP", f"{row['current_pcp_pct']:.1f}%")]
    tb = axt.table(cellText=kv, colLabels=["Metric", "Value"], cellLoc="left",
                   colWidths=[0.56, 0.44], bbox=[0, 0.02, 1, 0.96])
    tb.auto_set_font_size(False)
    tb.set_fontsize(7.5)
    for (r_, c_), cell in tb.get_celld().items():
        cell.set_edgecolor(C_GRID)
        if r_ == 0:
            cell.set_facecolor(C_HEADER)
            cell.set_text_props(color="white", weight="bold")
    axt.set_title("Key figures", loc="left")

    # (c) inventory sawtooth, one year
    ax = fig.add_subplot(gs[1, 0])
    T = q / it.D
    horizon = min(1.0, max(3 * T, 0.25))
    t = np.linspace(0, horizon, 2000)
    tau = np.mod(t, T)
    if it.m:
        tp = q / it.m
        inv = np.where(tau < tp, (it.m - it.D) * tau, q * it.k - it.D * (tau - tp))
    else:
        inv = q - it.D * tau
    ax.plot(t * dpy, inv, color=C_BLUE, lw=1.8, label="Recommended")
    if not pd.isna(row["current_q"]):
        qc = row["current_q"]
        Tc = qc / it.D
        tauc = np.mod(t, Tc)
        if it.m:
            tpc = qc / it.m
            invc = np.where(tauc < tpc, (it.m - it.D) * tauc, qc * it.k - it.D * (tauc - tpc))
        else:
            invc = qc - it.D * tauc
        ax.plot(t * dpy, invc, color=C_ORANGE, lw=1.2, alpha=0.8, label="Current")
    ax.axhline(q * it.k / 2, color=C_BLUE, ls="--", lw=0.9, label=f"Avg inventory {q * it.k / 2:,.0f}")
    if row["reorder_point"] > 0 and row["reorder_point"] < q * it.k:
        ax.axhline(row["reorder_point"], color=C_RED, ls=":", lw=1, label=f"Reorder point {row['reorder_point']:,.0f}")
    ax.set_xlabel("Working days")
    ax.set_ylabel("On-hand inventory (units)")
    ax.set_title("Inventory profile over time (sawtooth)", loc="left")
    ax.legend(fontsize=6.5, loc="upper right", ncol=1)
    ax.set_ylim(0, None)

    # (d) cost breakdown current vs recommended
    ax = fig.add_subplot(gs[1, 1])
    labels = ["Recommended"]
    ordc, hold = [row["annual_ordering_cost"]], [row["annual_holding_cost"]]
    disc_loss = [row["annual_purchase_cost"] - it.D * it.v0]
    if not pd.isna(row["current_q"]):
        cc = it.costs(row["current_q"])
        labels.insert(0, "Current")
        ordc.insert(0, cc["ordering"])
        hold.insert(0, cc["holding"])
        disc_loss.insert(0, cc["purchase"] - it.D * it.v0)
    x = np.arange(len(labels))
    ax.bar(x, ordc, 0.5, color=C_BLUE, label="Ordering/setup")
    ax.bar(x, hold, 0.5, bottom=ordc, color=C_ORANGE, label="Carrying")
    for xi, o, h in zip(x, ordc, hold):
        ax.text(xi, o + h, f"{cur}{o + h:,.0f}", ha="center", va="bottom", fontsize=8)
    labels = [lab + (f"\n(discount saves\n{cur}{-dl:,.0f}/yr on price)" if dl < -1e-6 else "")
              for lab, dl in zip(labels, disc_loss)]
    ax.set_xticks(x, labels, fontsize=7.5)
    ax.set_ylim(0, max(o + h for o, h in zip(ordc, hold)) * 1.35)
    ax.set_title("TRC breakdown (per year)", loc="left")
    ax.yaxis.set_major_formatter(_money(cur))
    ax.legend(fontsize=7, loc="upper center", ncol=2)
    ax.grid(axis="x", visible=False)
    return fig


def write_pdf(path, settings, res, items: List[Item], log, sku_pages: Optional[int] = None):
    kp = portfolio_kpis(res)
    cur = settings["currency"]
    by_sku = {it.sku: it for it in items}
    with PdfPages(path) as pdf:
        _cover(pdf, settings, kp)
        _text_page(pdf, "How to read this report", [
            ("What was calculated",
             "For every item the tool finds the order (or production) quantity Q that minimises the total annual "
             "cost = ordering/setup cost (A*D/Q) + carrying cost (Q/2 * v * r, multiplied by (1-D/m) for produced "
             "items) + purchase cost (D*v, which only matters when quantity discounts exist)."),
            ("Models used",
             "EOQ = sqrt(2AD / vr)  for purchased items (entire order arrives at once).\n"
             "EPQ = sqrt(2AD / (vr(1-D/m)))  for items produced in-house at rate m.\n"
             "All-units quantity discounts: the optimum of each price tier is found; if it falls below the tier's "
             "break quantity the break quantity is used; all tiers are then compared on total cost "
             "(lecture Cases 1, 2 and 3)."),
            ("Practical rules",
             "The theoretical optimum is rounded to the pack size and kept within the minimum and maximum order "
             "quantities. The column 'Practical rules applied' shows which rules shaped each answer."),
            ("Reorder point",
             "Place an order when the inventory position (on hand + on order) falls to D x L, where L is the "
             "lead time. Demand is assumed level and known; safety stock for uncertain demand is NOT included."),
            ("Flexibility band",
             f"Any order size between 'Flex band low' and 'Flex band high' costs at most "
             f"{settings['flex_tolerance_pct']:g}% more ordering+carrying cost than the optimum. The EOQ cost "
             "curve is flat near its minimum (PCP = 50 a^2/(1+a)), so rounding to pallets or truckloads inside "
             "this band is essentially free."),
            ("PCP of the current policy",
             "Percentage Cost Penalty = (cost of current Q - cost of recommended Q) / optimal TRC x 100."),
            ("Limitations",
             "Level, deterministic demand; costs constant over time; items treated independently (no joint "
             "replenishment); no shortages. Re-run the analysis when inputs change significantly."),
        ], subtitle=f"{settings['company_name']}  |  {kp['n_items']} items  |  currency {cur}")
        _portfolio_charts(pdf, settings, res)
        tbl = res[["sku", "model", "recommended_q", "textbook_q", "unit_price_paid", "orders_per_year",
                   "cycle_time_days", "reorder_point", "avg_inventory_value", "annual_relevant_cost",
                   "current_q", "annual_savings", "current_pcp_pct"]].copy()
        fmt = {"recommended_q": "{:,.0f}", "textbook_q": "{:,.1f}", "unit_price_paid": "{:,.3f}",
               "orders_per_year": "{:,.1f}", "cycle_time_days": "{:,.1f}", "reorder_point": "{:,.0f}",
               "avg_inventory_value": "{:,.0f}", "annual_relevant_cost": "{:,.0f}", "current_q": "{:,.0f}",
               "annual_savings": "{:,.0f}", "current_pcp_pct": "{:.1f}"}
        for c, f in fmt.items():
            tbl[c] = tbl[c].map(lambda x, f=f: "-" if pd.isna(x) else f.format(x))
        tbl.columns = ["SKU", "Model", "Rec. Q", "Textbook Q", f"Price ({cur})", "Orders/yr", "Cycle (days)",
                       "ROP", f"Avg inv ({cur})", f"TRC ({cur})", "Current Q", f"Savings ({cur})", "PCP %"]
        _table_pages(pdf, tbl, "Recommendations - summary table")
        if len(log):
            lg = log.copy()
            lg["message"] = lg["message"].map(lambda s: "\n".join(textwrap.wrap(str(s), 110)))
            _table_pages(pdf, lg, "Input validation log", col_widths=[0.12, 0.1, 0.78], rows_per_page=24)
        order = res.sort_values("annual_savings", ascending=False, na_position="last")
        if sku_pages is not None:
            order = order.head(sku_pages)
        for _, row in order.iterrows():
            fig = sku_figure(by_sku[row["sku"]], row, settings)
            pdf.savefig(fig)
            plt.close(fig)
        info = pdf.infodict()
        info["Title"] = settings["report_title"]
        info["Author"] = str(settings.get("analyst") or "Lot-Sizing Tool")


# =============================================================================
# 6. TEMPLATE WRITER
# =============================================================================
def write_input_workbook(path, settings: Optional[dict] = None, items: Optional[pd.DataFrame] = None,
                         breaks: Optional[pd.DataFrame] = None):
    """Create an input workbook (blank template when items is None) with instructions."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    st = {k: d for k, d, _ in SETTINGS_SPEC}
    st.update(settings or {})
    item_cols = [c for c, *_ in ITEM_COLUMNS]
    items = items if items is not None else pd.DataFrame(columns=item_cols)
    breaks = breaks if breaks is not None else pd.DataFrame(columns=["sku", "min_qty", "unit_cost"])
    instr = [
        ["LOT-SIZING TOOL - INPUT WORKBOOK"],
        [""],
        ["STEP 1  Fill the 'Settings' sheet (company name, currency, working days, default carrying rate r)."],
        ["STEP 2  Enter one row per item in the 'Items' sheet. Grey headers = optional, blue headers = required."],
        ["STEP 3  (optional) Enter supplier all-units price breaks in 'PriceBreaks' - one row per break."],
        ["STEP 4  Save the file and run:   python lot_sizing_tool.py --input <this file> --outdir output"],
        [""],
        ["RULES OF THUMB FOR THE DATA"],
        ["- Use the SAME time unit everywhere: annual_demand and production_rate per YEAR, carrying_rate per YEAR."],
        ["- order_cost A = cost that is incurred once per order/setup (PO processing, receiving, inspection, freight "
         "fixed charge, machine changeover). Do NOT include the price of the goods."],
        ["- carrying_rate r = cost of capital + storage + insurance + obsolescence, as a fraction of item value "
         "per year (typically 0.15 - 0.35). Enter 0.24, not 24."],
        ["- unit_cost v = value of one unit (purchase price or standard production cost) BEFORE discounts."],
        ["- Produced items need production_rate m (units/year while the line runs) greater than annual_demand."],
        ["- Price breaks are ALL-UNITS: when Q >= min_qty the whole order is charged the lower price."],
        ["- Leave optional cells blank rather than typing 0 (0 for max_order_qty means 'no limit')."],
        [""],
        ["See the 'Data_Dictionary' sheet for every column, its unit and meaning."],
    ]
    dd = pd.DataFrame([(c, "Items", "yes" if req else "no", t, u, d) for c, req, t, u, d in ITEM_COLUMNS] +
                      [(c, "PriceBreaks", "yes" if req else "no", t, u, d) for c, req, t, u, d in BREAK_COLUMNS] +
                      [(k, "Settings", "no", "", "", d) for k, _, d in SETTINGS_SPEC],
                      columns=["field", "sheet", "required", "type", "unit", "meaning"])
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame(instr).to_excel(xw, sheet_name="Instructions", index=False, header=False)
        pd.DataFrame([(k, st[k], d) for k, _, d in SETTINGS_SPEC], columns=["setting", "value", "description"]) \
            .to_excel(xw, sheet_name="Settings", index=False)
        items.reindex(columns=item_cols).to_excel(xw, sheet_name="Items", index=False)
        breaks.to_excel(xw, sheet_name="PriceBreaks", index=False)
        dd.to_excel(xw, sheet_name="Data_Dictionary", index=False)
    wb = load_workbook(path)
    req_fill, opt_fill = PatternFill("solid", fgColor="0F4C81"), PatternFill("solid", fgColor="7F8C8D")
    ws = wb["Instructions"]
    ws["A1"].font = Font(bold=True, size=14, color="0F4C81")
    ws["A8"].font = Font(bold=True)
    ws.column_dimensions["A"].width = 140
    required = {c for c, req, *_ in ITEM_COLUMNS if req} | {"sku", "min_qty", "unit_cost", "setting", "value"}
    for name in ("Settings", "Items", "PriceBreaks", "Data_Dictionary"):
        ws = wb[name]
        for cell in ws[1]:
            cell.fill = req_fill if (cell.value in required or name == "Data_Dictionary") else opt_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[cell.column_letter].width = 18
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
    wb["Settings"].column_dimensions["C"].width = 80
    wb["Data_Dictionary"].column_dimensions["F"].width = 95
    wb["Items"].column_dimensions["B"].width = 30
    dv = DataValidation(type="list", formula1='"Purchased,Produced"', allow_blank=False)
    wb["Items"].add_data_validation(dv)
    dv.add("D2:D2000")
    # comments on headers
    from openpyxl.comments import Comment
    for cell in wb["Items"][1]:
        for c, req, t, u, d in ITEM_COLUMNS:
            if cell.value == c:
                cell.comment = Comment(f"{d}\nUnit: {u}\nRequired: {'yes' if req else 'no'}", "Lot-Sizing Tool")
    wb.save(path)


# =============================================================================
# 7. ORCHESTRATION
# =============================================================================
def run(input_path: str, outdir: str = "output", pdf: bool = True, sku_pages: Optional[int] = None,
        verbose: bool = True) -> Dict:
    """Run the full pipeline. Returns dict with results DataFrame, KPIs, log and output paths."""
    os.makedirs(outdir, exist_ok=True)
    settings, items_df, breaks_df = load_input(input_path)
    items, log = build_items(settings, items_df, breaks_df)
    if not items:
        raise ValueError("No valid items to analyse - check the Validation log:\n" + log.to_string())
    res = analyse(settings, items)
    stem = os.path.splitext(os.path.basename(input_path))[0].replace("_input", "")
    xlsx = os.path.join(outdir, f"{stem}_results.xlsx")
    write_excel(xlsx, settings, res, log, items_df, breaks_df)
    pdf_path = None
    if pdf:
        pdf_path = os.path.join(outdir, f"{stem}_report.pdf")
        write_pdf(pdf_path, settings, res, items, log, sku_pages)
    kp = portfolio_kpis(res)
    if verbose:
        cur = settings["currency"]
        print(f"Lot-Sizing Tool v{__version__}  |  {settings['company_name']}")
        print(f"  items analysed        : {kp['n_items']}  (validation messages: {len(log)})")
        print(f"  annual savings        : {cur}{kp['total_savings']:,.2f}")
        print(f"  avg inventory value   : {cur}{kp['inv_value_current']:,.0f} (current) -> "
              f"{cur}{kp['inv_value_recommended']:,.0f} (recommended, all items)")
        print(f"  Excel results         : {xlsx}")
        if pdf_path:
            print(f"  PDF report            : {pdf_path}")
    return dict(results=res, kpis=kp, log=log, items=items, settings=settings, xlsx=xlsx, pdf=pdf_path)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Lot-sizing decision tool (EOQ / EPQ / quantity discounts)",
                                 formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    ap.add_argument("--input", "-i", help="Input workbook (.xlsx)")
    ap.add_argument("--outdir", "-o", default="output", help="Folder for results (default: output)")
    ap.add_argument("--no-pdf", action="store_true", help="Skip the PDF report")
    ap.add_argument("--sku-pages", type=int, default=None, help="Only the N SKUs with largest savings get a detail page")
    ap.add_argument("--template", metavar="PATH", help="Write a blank input template to PATH and exit")
    ap.add_argument("--version", action="version", version=__version__)
    a = ap.parse_args(argv)
    if a.template:
        write_input_workbook(a.template)
        print(f"Blank template written to {a.template}")
        return 0
    if not a.input:
        ap.print_help()
        return 1
    run(a.input, a.outdir, pdf=not a.no_pdf, sku_pages=a.sku_pages)
    return 0


if __name__ == "__main__":
    sys.exit(main())
